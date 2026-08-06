"""Residue-scan deterministic cases (plan Phase 7 / D8). Builds DOCX directly with
python-docx (no pandoc needed) so these run everywhere."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document

from mootloop.export.residue import scan_docx


def _docx(tmp_path: Path, text: str, name: str = "d.docx") -> Path:
    document = Document()
    document.add_paragraph(text)
    path = tmp_path / name
    document.save(str(path))
    return path


def test_clean_docx_passes(tmp_path: Path) -> None:
    path = _docx(tmp_path, "Defendant objects on relevance grounds. Admitted in part.")
    assert scan_docx(path).status == "pass"


def test_planted_annotation_marker_fails(tmp_path: Path) -> None:
    path = _docx(tmp_path, "associate: my self_assessment here with confidence=0.9")
    result = scan_docx(path)
    assert result.status == "fail"
    assert any(f.code == "annotation_marker" for f in result.findings)


def test_canary_marker_fails(tmp_path: Path) -> None:
    path = _docx(tmp_path, "Body text MOOTLOOP-CANARY leaked into the filing.")
    result = scan_docx(path)
    assert result.status == "fail"
    assert any(f.code == "annotation_marker" for f in result.findings)


def _with_comments_part(tmp_path: Path, comments_xml: str, name: str) -> Path:
    clean = _docx(tmp_path, "A clean paragraph.", name=f"src-{name}")
    doctored = tmp_path / name
    with zipfile.ZipFile(clean) as src, zipfile.ZipFile(doctored, "w") as dst:
        for item in src.namelist():
            dst.writestr(item, src.read(item))
        dst.writestr("word/comments.xml", comments_xml)
    return doctored


def test_comments_part_with_actual_comments_fails(tmp_path: Path) -> None:
    doctored = _with_comments_part(
        tmp_path,
        '<w:comments><w:comment w:id="1"><w:p>review note</w:p></w:comment></w:comments>',
        "with-comments.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(f.code == "comments_part" for f in result.findings)


def test_empty_comments_scaffold_passes(tmp_path: Path) -> None:
    # Pandoc >= 3.1 always emits an empty <w:comments/> part; an empty scaffold
    # carries no review-comment residue and must not block a clean export.
    doctored = _with_comments_part(tmp_path, "<w:comments/>", "empty-comments.docx")
    assert scan_docx(doctored).status == "pass"


def test_missing_file_fails(tmp_path: Path) -> None:
    result = scan_docx(tmp_path / "nope.docx")
    assert result.status == "fail"


# --- regression: parts the scan used to skip entirely ------------------------


def _with_part(tmp_path: Path, part: str, xml: str, name: str) -> Path:
    """A clean DOCX with one extra/overridden part spliced in."""
    clean = _docx(tmp_path, "A clean paragraph.", name=f"src-{name}")
    doctored = tmp_path / name
    with zipfile.ZipFile(clean) as src, zipfile.ZipFile(doctored, "w") as dst:
        for item in src.namelist():
            if item != part:
                dst.writestr(item, src.read(item))
        dst.writestr(part, xml)
    return doctored


def _body(runs: str) -> str:
    return (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        f'2006/main"><w:body><w:p>{runs}</w:p></w:body></w:document>'
    )


def test_marker_in_footnotes_fails(tmp_path: Path) -> None:
    # Pandoc renders markdown footnotes into word/footnotes.xml — a part the scan
    # used to ignore, so an attribution in a footnote shipped in a "clean" export.
    doctored = _with_part(
        tmp_path,
        "word/footnotes.xml",
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        '2006/main"><w:footnote w:id="2"><w:p><w:r><w:t>associate: internal note'
        "</w:t></w:r></w:p></w:footnote></w:footnotes>",
        "footnote-marker.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(
        f.code == "annotation_marker" and "footnotes" in (f.locator or "")
        for f in result.findings
    )


def test_marker_in_endnotes_fails(tmp_path: Path) -> None:
    doctored = _with_part(
        tmp_path,
        "word/endnotes.xml",
        '<w:endnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        '2006/main"><w:endnote w:id="2"><w:p><w:r><w:t>MOOTLOOP-CANARY</w:t></w:r>'
        "</w:p></w:endnote></w:endnotes>",
        "endnote-marker.docx",
    )
    assert scan_docx(doctored).status == "fail"


def test_marker_in_core_properties_fails(tmp_path: Path) -> None:
    doctored = _with_part(
        tmp_path,
        "docProps/core.xml",
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/'
        'metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:description>MOOTLOOP-CANARY</dc:description></cp:coreProperties>",
        "meta-marker.docx",
    )
    assert scan_docx(doctored).status == "fail"


def test_marker_split_across_runs_fails(tmp_path: Path) -> None:
    # A Word round-trip splits a run at rsid boundaries; the marker only reassembles
    # in the extracted visible text, never in a raw-substring scan.
    doctored = _with_part(
        tmp_path,
        "word/document.xml",
        _body("<w:r><w:t>associ</w:t></w:r><w:r><w:t>ate: my note</w:t></w:r>"),
        "split-marker.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(f.code == "annotation_marker" for f in result.findings)


# --- regression: tracked changes are matched as ELEMENTS, not substrings -----


def test_table_inside_borders_do_not_read_as_tracked_changes(tmp_path: Path) -> None:
    # `<w:insideH>`/`<w:insideV>` are ordinary table borders. A `"<w:ins" in raw`
    # test reads them as an insertion and deletes a perfectly clean filing.
    doctored = _with_part(
        tmp_path,
        "word/document.xml",
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        '2006/main"><w:body><w:tbl><w:tblPr><w:tblBorders>'
        '<w:insideH w:val="single"/><w:insideV w:val="single"/>'
        "</w:tblBorders></w:tblPr></w:tbl></w:body></w:document>",
        "table-borders.docx",
    )
    assert scan_docx(doctored).status == "pass"


def test_real_insertion_still_fails(tmp_path: Path) -> None:
    doctored = _with_part(
        tmp_path,
        "word/document.xml",
        _body('<w:ins w:id="1" w:author="Jane"><w:r><w:t>added</w:t></w:r></w:ins>'),
        "tracked-ins.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(f.code == "tracked_changes" for f in result.findings)


def test_formatting_revision_fails(tmp_path: Path) -> None:
    # `w:rPrChange` carries the reviser's name and timestamp into the filed copy.
    doctored = _with_part(
        tmp_path,
        "word/document.xml",
        _body(
            '<w:r><w:rPr><w:rPrChange w:id="7" w:author="Jane" w:date="2026-01-01T00:00:00Z">'
            "<w:rPr/></w:rPrChange></w:rPr><w:t>text</w:t></w:r>"
        ),
        "tracked-format.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(f.code == "tracked_changes" for f in result.findings)


def test_dangling_comment_anchor_fails(tmp_path: Path) -> None:
    # The comments part can be stripped while the in-body anchors survive.
    doctored = _with_part(
        tmp_path,
        "word/document.xml",
        _body('<w:commentRangeStart w:id="1"/><w:r><w:t>text</w:t></w:r>'),
        "dangling-anchor.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(f.code == "comments_part" for f in result.findings)


def test_tracked_change_in_footnotes_fails(tmp_path: Path) -> None:
    doctored = _with_part(
        tmp_path,
        "word/footnotes.xml",
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        '2006/main"><w:footnote w:id="2"><w:p>'
        '<w:del w:id="3" w:author="Jane"><w:r><w:delText>cut</w:delText></w:r></w:del>'
        "</w:p></w:footnote></w:footnotes>",
        "footnote-tracked.docx",
    )
    result = scan_docx(doctored)
    assert result.status == "fail"
    assert any(f.code == "tracked_changes" for f in result.findings)


# --- regression: draft-only renderer placeholders ----------------------------


def test_unsigned_signature_placeholder_fails(tmp_path: Path) -> None:
    # `master.py` renders `[ATTORNEY NAME]` when the matter has no attorney block.
    path = _docx(tmp_path, "Respectfully submitted,\n[ATTORNEY NAME]", name="unsigned.docx")
    result = scan_docx(path)
    assert result.status == "fail"
    assert any(f.code == "annotation_marker" for f in result.findings)


def test_undrafted_response_placeholder_fails(tmp_path: Path) -> None:
    # `_no response drafted_` loses its markdown emphasis in the DOCX — match the
    # rendered text, which is what actually reaches the page.
    path = _docx(tmp_path, "RESPONSE: no response drafted", name="undrafted.docx")
    assert scan_docx(path).status == "fail"


def test_unnamed_party_placeholder_fails(tmp_path: Path) -> None:
    path = _docx(tmp_path, "[Plaintiff],\n\nv.\n\n[Defendant],", name="unnamed.docx")
    assert scan_docx(path).status == "fail"


def test_real_caption_still_passes(tmp_path: Path) -> None:
    # The un-bracketed caption a real matter renders must not trip the placeholders.
    path = _docx(
        tmp_path,
        "Northfield Widgets LLC,\n\nPlaintiff,\n\nv.\n\nGranite Supply Co.,\n\nDefendant.",
        name="real-caption.docx",
    )
    assert scan_docx(path).status == "pass"
